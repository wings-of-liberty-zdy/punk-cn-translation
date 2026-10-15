#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把翻译表导出成方便批量编辑的 Excel。

产出 translations/punk-zh-CN.xlsx，两张表：
  「翻译表」  371 行，只改「译文」列，其余列是参考信息
  「说明」    编辑规则、怎么改回来

配套 tools/import_xlsx.py 把改好的 Excel 导回流水线。

用法：
    python tools/export_xlsx.py
    python tools/export_xlsx.py --out translations/punk-zh-CN.xlsx
"""

import argparse
import collections
import json
import os
import re
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABLE = os.path.join(ROOT, "translations", "punk.tsv")
REPORT = os.path.join(ROOT, "data", "validate.json")
CLASSIFIED = os.path.join(ROOT, "data", "dump", "strings_all_classified.jsonl")
OUT = os.path.join(ROOT, "translations", "punk-zh-CN.xlsx")

CAT_CN = {
    "ui_menu": "界面与菜单",
    "item_module_desc": "物品与模块",
    "tutorial_hint": "教程提示",
    "dialogue_story": "剧情",
    "other": "其他",
}
_QUOTES = {"\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"'}


def norm(s):
    for a, b in _QUOTES.items():
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s).strip()


def load_contexts():
    ctx = {}
    if not os.path.exists(CLASSIFIED):
        return ctx
    for line in open(CLASSIFIED, encoding="utf-8"):
        line = line.strip()
        if not line:
            continue
        o = json.loads(line)
        if o.get("decision") != "keep":
            continue
        k = norm(o["text"])
        if k not in ctx:
            ctx[k] = o.get("key_or_context", "")
    return ctx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--table", default=TABLE)
    ap.add_argument("--out", default=OUT)
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    rows = []
    for line in open(args.table, encoding="utf-8"):
        line = line.rstrip("\n")
        if not line or line.startswith("#"):
            continue
        f = line.split("\t")
        if len(f) < 3:
            continue
        rows.append({"id": f[0], "en": f[1], "cn": f[2]})

    ctxs = load_contexts()
    cats, flags = {}, {}
    if os.path.exists(CLASSIFIED):
        for line in open(CLASSIFIED, encoding="utf-8"):
            line = line.strip()
            if not line:
                continue
            o = json.loads(line)
            if o.get("decision") == "keep":
                cats.setdefault(norm(o["text"]), o.get("category", "other"))
    if os.path.exists(REPORT):
        rep = json.load(open(REPORT, encoding="utf-8"))
        for rid, lst in rep.get("issues", {}).items():
            flags[rid] = "；".join("%s：%s" % (d["kind"], d["detail"]) if d["detail"] else d["kind"]
                                   for d in lst)

    wb = Workbook()
    ws = wb.active
    ws.title = "翻译表"

    head = ["编号", "分类", "原文（英文）", "译文（中文）", "校验提示", "出现位置"]
    widths = [8, 12, 52, 52, 30, 60]
    ws.append(head)

    hfill = PatternFill("solid", fgColor="1F3864")
    hfont = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (h, w) in enumerate(zip(head, widths), start=1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = hfill, hfont
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    cn_font = Font(size=11)
    for r in rows:
        k = norm(r["en"])
        ws.append([r["id"],
                   CAT_CN.get(cats.get(k, "other"), cats.get(k, "other")),
                   r["en"],
                   r["cn"],
                   flags.get(r["id"], ""),
                   ctxs.get(k, "")])
        i = ws.max_row
        for j in range(1, 7):
            cell = ws.cell(row=i, column=j)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(j in (3, 4, 5, 6)))
        ws.cell(row=i, column=4).font = cn_font
        if flags.get(r["id"]):
            ws.cell(row=i, column=5).fill = warn_fill
    ws.auto_filter.ref = "A1:F%d" % ws.max_row

    # ── 说明表 ──
    ws2 = wb.create_sheet("说明")
    ws2.column_dimensions["A"].width = 100
    notes = [
        ("PUNK 汉化 · 翻译表（%d 条）" % len(rows), True),
        ("", False),
        ("怎么用", True),
        ("  只改「译文（中文）」这一列（D 列）。其他列是参考信息，改了也不会生效。", False),
        ("  改完保存，运行：python tools\\import_xlsx.py", False),
        ("  脚本会自动和当前翻译表比对，把改动写成 translations\\overlay\\manual-edits.tsv", False),
        ("  然后运行 python tools\\build_table.py --sync-plugin 生效。", False),
        ("", False),
        ("换行怎么写（重要）", True),
        ("  单元格里的换行写成字面量 \\n （反斜杠 + 字母 n），不要在单元格里按 Alt+Enter 输入真换行。", False),
        ("  例：选择\\n配装   表示界面上分两行显示。", False),
        ("  规则：译文的 \\n 数量不得超过原文。少是可以的，中文短，合成一行更安全。", False),
        ("", False),
        ("长度怎么把握（重要）", True),
        ("  宽度按渲染算：汉字 = 2，空格 = 0.5，其余字符 = 1。译文宽度不要超过原文。", False),
        ("  不要按字符数算 —— 一个汉字约占两个拉丁字母宽。按钮和标题栏尤其要小心。", False),
        ("  「校验提示」列标出的条目建议重点看，但不是硬错误。", False),
        ("", False),
        ("必须原样保留的东西", True),
        ("  富文本标签：<color=#RRGGBB> </color> <sprite=5> 等，数量和配对都不能变。", False),
        ("  占位符：{0} {1} {0:hh':'mm':'ss} 等，原样保留。", False),
        ("  HUD 资源名保留英文大写：CAPS ELECTRON FUEL GEL TECH STAMINA HEALTH POWER", False),
        ("  按键名保留英文：TAB ALT SPACE ENTER ESC SHIFT CTRL WASD 等。", False),
        ("  原文里的拼写错误也照抄（例如 exlode、pontential），那是原作者的笔误，不要顺手改。", False),
        ("", False),
        ("标点", True),
        ("  沿用原文的半角标点（, . ? :），不要换成全角（，。？：）。", False),
        ("", False),
        ("语气", True),
        ("  冷峻科幻腔。少用感叹号和语气词。术语见 glossary/terms.tsv。", False),
        ("", False),
        ("分类图例", True),
        ("  界面与菜单 / 物品与模块 / 教程提示 / 剧情 / 其他", False),
    ]
    for text, bold in notes:
        ws2.append([text])
        ws2.cell(row=ws2.max_row, column=1).font = Font(bold=bold, size=12 if bold else 11)
        ws2.cell(row=ws2.max_row, column=1).alignment = Alignment(vertical="center", wrap_text=False)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)

    counts = collections.Counter()
    for r in rows:
        counts[CAT_CN.get(cats.get(norm(r["en"]), "other"), "其他")] += 1
    print("已写出 %s" % os.path.relpath(args.out, ROOT))
    print("  翻译表 %d 行，其中带校验提示 %d 行" % (len(rows), sum(1 for r in rows if flags.get(r["id"]))))
    print("  分类：" + "、".join("%s %d" % (k, v) for k, v in counts.most_common()))
    print("  改完运行：python tools\\import_xlsx.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
